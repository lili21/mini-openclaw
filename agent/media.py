import base64
import logging
import os

from lark_oapi.api.im.v1 import GetFileRequest, GetImageRequest

from config import MAX_FILE_SIZE, UPLOAD_DIR

logger = logging.getLogger(__name__)


def download_feishu_image(client, image_key: str) -> tuple[bytes | None, str | None]:
    request = GetImageRequest.builder().image_key(image_key).build()
    response = client.im.v1.image.get(request)

    if not response.success():
        logger.error(
            f"Failed to download image: code={response.code}, msg={response.msg}"
        )
        return None, None

    image_data = response.data.image

    ext = _detect_image_extension(image_data)
    filename = f"{image_key}.{ext}"
    filepath = os.path.join(UPLOAD_DIR, filename)

    with open(filepath, "wb") as f:
        f.write(image_data)

    logger.info(f"Image saved to {filepath}")
    return image_data, filepath


def download_feishu_file(
    client, file_key: str, filename: str
) -> tuple[str | None, int | None]:
    request = GetFileRequest.builder().file_key(file_key).build()
    response = client.im.v1.file.get(request)

    if not response.success():
        logger.error(
            f"Failed to download file: code={response.code}, msg={response.msg}"
        )
        return None, None

    file_data = response.data.file
    file_size = len(file_data)

    if file_size > MAX_FILE_SIZE:
        logger.warning(f"File too large: {file_size} > {MAX_FILE_SIZE}")
        return None, file_size

    filepath = os.path.join(UPLOAD_DIR, filename)
    with open(filepath, "wb") as f:
        f.write(file_data)

    logger.info(f"File saved to {filepath}, size: {file_size}")
    return filepath, file_size


def image_to_base64_url(image_data: bytes) -> str:
    ext = _detect_image_extension(image_data)
    b64_data = base64.b64encode(image_data).decode("utf-8")
    return f"data:image/{ext};base64,{b64_data}"


def parse_file_content(filepath: str) -> str | None:
    ext = os.path.splitext(filepath)[1].lower()

    try:
        if ext == ".pdf":
            return _parse_pdf(filepath)
        elif ext in [".docx", ".doc"]:
            return _parse_docx(filepath)
        elif ext in [".xlsx", ".xls"]:
            return _parse_excel(filepath)
        elif ext in [
            ".txt",
            ".md",
            ".json",
            ".csv",
            ".py",
            ".js",
            ".ts",
            ".html",
            ".css",
        ]:
            return _parse_text(filepath)
        else:
            return f"不支持的文件格式: {ext}"
    except Exception as e:
        logger.error(f"Failed to parse file {filepath}: {e}")
        return f"文件解析失败: {str(e)}"


def _detect_image_extension(data: bytes) -> str:
    if data[:8] == b"\x89PNG\r\n\x1a\n":
        return "png"
    elif data[:2] == b"\xff\xd8":
        return "jpeg"
    elif data[:6] in (b"GIF87a", b"GIF89a"):
        return "gif"
    elif data[:4] == b"RIFF" and data[8:12] == b"WEBP":
        return "webp"
    return "png"


def _parse_pdf(filepath: str) -> str:
    from pypdf import PdfReader

    reader = PdfReader(filepath)
    text_parts = []
    for page in reader.pages:
        text = page.extract_text()
        if text:
            text_parts.append(text)

    return "\n\n".join(text_parts)


def _parse_docx(filepath: str) -> str:
    from docx import Document

    doc = Document(filepath)
    text_parts = [para.text for para in doc.paragraphs if para.text]
    return "\n".join(text_parts)


def _parse_excel(filepath: str) -> str:
    from openpyxl import load_workbook

    wb = load_workbook(filepath, data_only=True)
    text_parts = []

    for sheet_name in wb.sheetnames:
        sheet = wb[sheet_name]
        text_parts.append(f"=== Sheet: {sheet_name} ===")
        for row in sheet.iter_rows(values_only=True):
            row_text = " | ".join(str(cell) if cell is not None else "" for cell in row)
            if row_text.strip():
                text_parts.append(row_text)

    return "\n".join(text_parts)


def _parse_text(filepath: str) -> str:
    with open(filepath, "r", encoding="utf-8") as f:
        return f.read()
